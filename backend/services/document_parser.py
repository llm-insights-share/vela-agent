import io
import re
from typing import Any, Callable, Dict, List, Optional, Tuple


def parse_pdf(file_bytes: bytes) -> str:
    """Legacy whole-document text extract (knowledge base / flat consumers)."""
    units = parse_pdf_units(file_bytes)
    parts = [(u.get("text") or "").strip() for u in units if (u.get("text") or "").strip()]
    return "\n\n".join(parts)


# Soft limit for eager OCR on upload; remaining pages stay pending until asked.
MAX_EAGER_OCR_PAGES = 20
MIN_TEXT_LAYER_CHARS = 30
MAX_PAGE_IMAGE_SIDE = 1600
# Pages beyond this skip rasterization / vision (text layer only or pending without image).
MAX_VISION_PDF_PAGES = 100
# Kept for callers that still reference the old constant.
MAX_SCANNED_PDF_PAGES = 12


def _guess_image_type(name: str, data: bytes) -> Tuple[str, str]:
    n = (name or "").lower()
    if data[:2] == b"\xff\xd8" or n.endswith(".jpg") or n.endswith(".jpeg"):
        return ".jpg", "image/jpeg"
    if data[:8] == b"\x89PNG\r\n\x1a\n" or n.endswith(".png"):
        return ".png", "image/png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return ".webp", "image/webp"
    if n.endswith(".gif") or data[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif", "image/gif"
    return ".png", "image/png"


def _downscale_image(data: bytes, ext: str, mime: str) -> Tuple[bytes, str, str]:
    from PIL import Image
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
    except Exception:
        return data, ext, mime
    width, height = image.size
    if max(width, height) <= MAX_PAGE_IMAGE_SIDE:
        return data, ext, mime
    ratio = MAX_PAGE_IMAGE_SIDE / float(max(width, height))
    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")
    image = image.resize(
        (max(1, int(width * ratio)), max(1, int(height * ratio))),
        Image.Resampling.LANCZOS,
    )
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=85)
    return buf.getvalue(), ".jpg", "image/jpeg"


def _extract_xobject_raw(page) -> List[Tuple[int, bytes, str, str]]:
    results: List[Tuple[int, bytes, str, str]] = []
    resources = page.get("/Resources")
    if resources is None:
        return results
    resources = resources.get_object() if hasattr(resources, "get_object") else resources
    xobj = resources.get("/XObject") if resources else None
    if xobj is None:
        return results
    xobj = xobj.get_object() if hasattr(xobj, "get_object") else xobj
    for name in xobj:
        obj = xobj[name]
        obj = obj.get_object() if hasattr(obj, "get_object") else obj
        if obj.get("/Subtype") != "/Image":
            continue
        filt = obj.get("/Filter")
        filters = [str(f) for f in (filt if isinstance(filt, list) else ([filt] if filt else []))]
        try:
            data = obj.get_data()
        except Exception:
            continue
        if not data:
            continue
        if "/DCTDecode" in filters:
            results.append((len(data), data, ".jpg", "image/jpeg"))
            continue
        try:
            from PIL import Image
            width = int(obj["/Width"])
            height = int(obj["/Height"])
            cs = str(obj.get("/ColorSpace"))
            if "RGB" in cs:
                mode = "RGB"
            elif "Gray" in cs:
                mode = "L"
            elif "CMYK" in cs:
                mode = "CMYK"
            else:
                mode = "RGB" if len(data) >= width * height * 3 else "L"
            img = Image.frombytes(mode, (width, height), data)
            if img.mode == "CMYK":
                img = img.convert("RGB")
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            blob = buf.getvalue()
            results.append((len(blob), blob, ".png", "image/png"))
        except Exception:
            continue
    return results


def _extract_embedded_page_image(page) -> Optional[Tuple[bytes, str, str]]:
    """Pick the largest embedded image on a PDF page, if any."""
    candidates: List[Tuple[int, bytes, str, str]] = []
    try:
        for image in page.images:
            data = getattr(image, "data", None)
            if not data:
                continue
            ext, mime = _guess_image_type(getattr(image, "name", ""), data)
            candidates.append((len(data), data, ext, mime))
    except Exception:
        candidates = []
    if not candidates:
        try:
            candidates = _extract_xobject_raw(page)
        except Exception:
            candidates = []
    if not candidates:
        return None
    _, data, ext, mime = max(candidates, key=lambda item: item[0])
    data, ext, mime = _downscale_image(data, ext, mime)
    return data, ext, mime


def _render_page_with_pypdfium2(file_bytes: bytes, page_index: int) -> Optional[Tuple[bytes, str, str]]:
    """Rasterize a PDF page (0-based) via pypdfium2. Returns (bytes, ext, mime)."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return None
    try:
        pdf = pdfium.PdfDocument(file_bytes)
        try:
            if page_index < 0 or page_index >= len(pdf):
                return None
            page = pdf[page_index]
            # Aim for ~150 DPI; scale so longest side stays near MAX_PAGE_IMAGE_SIDE.
            width = max(page.get_width(), 1)
            height = max(page.get_height(), 1)
            scale = min(2.0, MAX_PAGE_IMAGE_SIDE / float(max(width, height)))
            bitmap = page.render(scale=scale)
            pil = bitmap.to_pil()
            if pil.mode not in ("RGB", "L"):
                pil = pil.convert("RGB")
            buf = io.BytesIO()
            pil.save(buf, format="JPEG", quality=85)
            return buf.getvalue(), ".jpg", "image/jpeg"
        finally:
            pdf.close()
    except Exception:
        return None


def extract_pdf_page_images(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Legacy: extract one image per page (capped). Prefer parse_pdf_units."""
    units = parse_pdf_units(file_bytes)
    out: List[Dict[str, Any]] = []
    for u in units[:MAX_SCANNED_PDF_PAGES]:
        raw = u.get("_image_bytes")
        if not raw:
            continue
        out.append({
            "page": u.get("unit_index"),
            "bytes": raw,
            "ext": u.get("_image_ext") or ".jpg",
            "mime": u.get("_image_mime") or "image/jpeg",
        })
    return out


def parse_pdf_units(file_bytes: bytes) -> List[Dict[str, Any]]:
    """
    Split a PDF into per-page document units.

    Each unit has:
      unit_index, text, source, status,
      and optionally _image_bytes/_image_ext/_image_mime (ephemeral, not persisted).
    """
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = list(reader.pages)
    units: List[Dict[str, Any]] = []

    for idx, page in enumerate(pages):
        page_no = idx + 1
        try:
            raw_text = (page.extract_text() or "").strip()
        except Exception:
            raw_text = ""

        has_text = len(raw_text) >= MIN_TEXT_LAYER_CHARS
        image_tuple = None
        # Long PDFs: only first MAX_VISION_PDF_PAGES get raster/embed images
        allow_vision = page_no <= MAX_VISION_PDF_PAGES
        if not has_text and allow_vision:
            image_tuple = _extract_embedded_page_image(page)
            if image_tuple is None:
                image_tuple = _render_page_with_pypdfium2(file_bytes, idx)

        if has_text and image_tuple is None:
            # Text-rich page; still try embedding if mixed content is useful later — skip for now.
            units.append({
                "unit_index": page_no,
                "text": raw_text,
                "source": "text_layer",
                "status": "ready",
            })
        elif has_text and image_tuple is not None:
            data, ext, mime = image_tuple
            units.append({
                "unit_index": page_no,
                "text": raw_text,
                "source": "mixed",
                "status": "ready",
                "_image_bytes": data,
                "_image_ext": ext,
                "_image_mime": mime,
            })
        elif image_tuple is not None:
            data, ext, mime = image_tuple
            units.append({
                "unit_index": page_no,
                "text": "",
                "source": "ocr",
                "status": "pending",
                "_image_bytes": data,
                "_image_ext": ext,
                "_image_mime": mime,
            })
        elif raw_text:
            # Short text layer, no image — keep whatever we have.
            units.append({
                "unit_index": page_no,
                "text": raw_text,
                "source": "text_layer",
                "status": "ready",
            })
        else:
            units.append({
                "unit_index": page_no,
                "text": "",
                "source": "ocr",
                "status": "failed",
            })

    return units


def parse_docx(file_bytes: bytes) -> str:
    units = parse_docx_units(file_bytes)
    return "\n\n".join((u.get("text") or "").strip() for u in units if (u.get("text") or "").strip())


def parse_docx_units(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Paragraphs and tables as ordered units."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(file_bytes))
    units: List[Dict[str, Any]] = []
    idx = 0
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            # Find matching paragraph
            text = ""
            for para in doc.paragraphs:
                if para._element is child:
                    text = (para.text or "").strip()
                    break
            if text:
                idx += 1
                units.append({
                    "unit_index": idx,
                    "text": text,
                    "source": "text_layer",
                    "status": "ready",
                })
        elif tag == qn("w:tbl"):
            rows_out: List[str] = []
            for table in doc.tables:
                if table._element is child:
                    for row in table.rows:
                        cells = []
                        for cell in row.cells:
                            cell_text = (cell.text or "").strip().replace("\n", " ")
                            if cell_text:
                                cells.append(cell_text)
                        if cells:
                            rows_out.append("\t".join(cells))
                    break
            if rows_out:
                idx += 1
                units.append({
                    "unit_index": idx,
                    "text": f"[表格]\n" + "\n".join(rows_out),
                    "source": "text_layer",
                    "status": "ready",
                })
    if not units:
        # Fallback: paragraphs only
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if text:
                idx += 1
                units.append({
                    "unit_index": idx,
                    "text": text,
                    "source": "text_layer",
                    "status": "ready",
                })
    return units


def parse_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace")


def parse_markdown(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"```[\s\S]*?```", "", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"[*_~>|]", "", text)
    return text


def parse_csv(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8-sig", errors="replace")
    return text.strip()


def parse_html(file_bytes: bytes) -> str:
    raw = file_bytes.decode("utf-8", errors="replace")
    try:
        from readability import Document as ReadabilityDocument
        doc = ReadabilityDocument(raw)
        html = doc.summary() or raw
    except Exception:
        html = raw
    # Strip tags roughly
    text = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html)
    text = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_xlsx(file_bytes: bytes) -> str:
    units = parse_xlsx_units(file_bytes)
    return "\n\n".join((u.get("text") or "") for u in units if u.get("text"))


def parse_xlsx_units(file_bytes: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    units: List[Dict[str, Any]] = []
    for sheet_idx, sheet in enumerate(wb.worksheets, start=1):
        lines = [f"[Sheet: {sheet.title}]"]
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append("\t".join(cells))
        text = "\n".join(lines)
        if len(lines) > 1:
            units.append({
                "unit_index": sheet_idx,
                "text": text,
                "source": "text_layer",
                "status": "ready",
            })
    wb.close()
    return units


def parse_xls(file_bytes: bytes) -> str:
    units = parse_xls_units(file_bytes)
    return "\n\n".join((u.get("text") or "") for u in units if u.get("text"))


def parse_xls_units(file_bytes: bytes) -> List[Dict[str, Any]]:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=file_bytes)
        units: List[Dict[str, Any]] = []
        for sheet_idx in range(book.nsheets):
            sheet = book.sheet_by_index(sheet_idx)
            lines = [f"[Sheet: {sheet.name}]"]
            for row_idx in range(sheet.nrows):
                cells = []
                for col_idx in range(sheet.ncols):
                    val = sheet.cell_value(row_idx, col_idx)
                    if val is None or val == "":
                        continue
                    text = str(val).strip()
                    if text:
                        cells.append(text)
                if cells:
                    lines.append("\t".join(cells))
            if len(lines) > 1:
                units.append({
                    "unit_index": sheet_idx + 1,
                    "text": "\n".join(lines),
                    "source": "text_layer",
                    "status": "ready",
                })
        return units
    except Exception as e:
        raise ValueError(f"无法解析 .xls 文件: {e}") from e


def parse_image(file_bytes: bytes) -> str:
    return ""


def _reject_legacy_doc(file_bytes: bytes) -> str:
    raise ValueError(
        "不支持旧版 .doc（二进制 Word）。请另存为 .docx 后再上传。"
    )


PARSERS: Dict[str, Callable[[bytes], str]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".doc": _reject_legacy_doc,
    ".txt": parse_txt,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".csv": parse_csv,
    ".html": parse_html,
    ".htm": parse_html,
    ".xlsx": parse_xlsx,
    ".xls": parse_xls,
    ".png": parse_image,
    ".jpg": parse_image,
    ".jpeg": parse_image,
    ".webp": parse_image,
    ".gif": parse_image,
}

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}

SUPPORTED_EXTENSIONS = sorted(PARSERS.keys())


def chunk_by_paragraph(text: str, target_size: int = 500, min_size: int = 100) -> List[str]:
    raw_paragraphs = re.split(r"\n\s*\n", text)
    paragraphs = [p.strip() for p in raw_paragraphs if p.strip()]
    if not paragraphs:
        return []

    chunks = []
    current = ""

    for para in paragraphs:
        if len(para) >= target_size:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(para):
                end = start + target_size
                chunks.append(para[start:end].strip())
                start = end
        elif current and len(current) + len(para) + 2 > target_size:
            chunks.append(current.strip())
            current = para
        else:
            if current:
                current += "\n\n" + para
            else:
                current = para

    if current:
        chunks.append(current.strip())

    result = []
    for chunk in chunks:
        if len(chunk) >= min_size:
            result.append(chunk)

    return result


def chunks_from_units(units: List[Dict[str, Any]], target_size: int = 500) -> List[str]:
    """Build BM25 chunks prefixed with [p.N] from ready units."""
    chunks: List[str] = []
    for u in units or []:
        text = (u.get("text") or "").strip()
        if not text or u.get("status") == "failed":
            continue
        idx = u.get("unit_index")
        prefix = f"[p.{idx}] " if idx is not None else ""
        paras = chunk_by_paragraph(text, target_size=target_size, min_size=1)
        if not paras:
            chunks.append(prefix + text)
        else:
            for p in paras:
                chunks.append(prefix + p)
    return chunks


def parsed_text_from_units(units: List[Dict[str, Any]], filename: str = "") -> str:
    """Concatenate ready unit texts for legacy parsed_text field."""
    parts = []
    for u in units or []:
        text = (u.get("text") or "").strip()
        if not text:
            continue
        idx = u.get("unit_index")
        if idx is not None and len(units) > 1:
            parts.append(f"[第{idx}页]\n{text}")
        else:
            parts.append(text)
    if parts:
        return "\n\n".join(parts)
    if any(u.get("status") == "pending" for u in (units or [])):
        return f"[扫描件] {filename}" if filename else "[扫描件]"
    return ""


def parse_file(
    file_bytes: bytes,
    filename: str,
) -> Tuple[str, bool, List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Parse file bytes.

    Returns:
        (text, is_image, page_images, units)

    page_images entries: {page, bytes, ext, mime}
    units: document units with optional ephemeral _image_* fields.
    """
    import os
    ext = os.path.splitext(filename)[1].lower()
    parser = PARSERS.get(ext)
    if parser is None:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(f"不支持的文件格式 {ext or '(无扩展名)'}，支持: {supported}")

    is_image = ext in IMAGE_EXTS
    page_images: List[Dict[str, Any]] = []
    units: List[Dict[str, Any]] = []

    if is_image:
        units = [{
            "unit_index": 1,
            "text": "",
            "source": "ocr",
            "status": "pending",
            "_image_bytes": file_bytes,
            "_image_ext": ext or ".png",
            "_image_mime": {
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
                ".webp": "image/webp",
                ".gif": "image/gif",
            }.get(ext, "image/png"),
        }]
        text = f"[图片] {filename}"
        return text, True, page_images, units

    if ext == ".pdf":
        units = parse_pdf_units(file_bytes)
        for u in units:
            raw = u.get("_image_bytes")
            if raw:
                page_images.append({
                    "page": u.get("unit_index"),
                    "bytes": raw,
                    "ext": u.get("_image_ext") or ".jpg",
                    "mime": u.get("_image_mime") or "image/jpeg",
                })
        text = parsed_text_from_units(units, filename)
        if not text.strip() and not page_images and not any(
            (u.get("text") or "").strip() for u in units
        ):
            if not units:
                raise ValueError("文件内容为空或无法解析")
            if all(u.get("status") == "failed" for u in units):
                raise ValueError("文件内容为空或无法解析")
        if not text.strip() and page_images:
            text = f"[扫描件] {filename}"
        if len(units) > MAX_VISION_PDF_PAGES:
            text = (text or "") + (
                f"\n\n[提示] 文档共 {len(units)} 页，"
                f"仅前 {MAX_VISION_PDF_PAGES} 页支持视觉/扫描解析；"
                "其余页仅在有文字层时可检索。"
            )
        return text, False, page_images, units

    if ext == ".docx":
        units = parse_docx_units(file_bytes)
        text = parsed_text_from_units(units, filename)
        if not text.strip():
            raise ValueError("文件内容为空或无法解析")
        return text, False, page_images, units

    if ext == ".xlsx":
        units = parse_xlsx_units(file_bytes)
        text = parsed_text_from_units(units, filename)
        if not text.strip():
            raise ValueError("文件内容为空或无法解析")
        return text, False, page_images, units

    if ext == ".xls":
        units = parse_xls_units(file_bytes)
        text = parsed_text_from_units(units, filename)
        if not text.strip():
            raise ValueError("文件内容为空或无法解析")
        return text, False, page_images, units

    # Other document types → single ready unit (.doc raises via parser)
    text = parser(file_bytes)
    if not text or not text.strip():
        raise ValueError("文件内容为空或无法解析")
    units = [{
        "unit_index": 1,
        "text": text,
        "source": "text_layer",
        "status": "ready",
    }]
    return text, False, page_images, units
